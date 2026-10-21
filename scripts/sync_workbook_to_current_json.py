#!/usr/bin/env python3
"""Synchronise the formatted workbook with the current 55-L3 JSON master."""

from datetime import datetime
import json
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
SOURCE = DATA / "heart_l4_risk_cards.json"
WORKBOOK = DATA / "HEART_L4_Risk_Cards.xlsx"
TEMP = DATA / ".HEART_L4_Risk_Cards.sync.xlsx"


def excel_value(key: str, value):
    if value in (None, ""):
        return None
    if key == "Evidence_Accessed_At":
        return datetime.fromisoformat(value).replace(tzinfo=None)
    return value


def main() -> None:
    cards = json.loads(SOURCE.read_text(encoding="utf-8"))
    assert len(cards) == 622
    assert len({card["L4_ID"] for card in cards}) == 622
    assert len({card["L3_ID"] for card in cards}) == 55

    workbook = openpyxl.load_workbook(WORKBOOK)
    sheet = workbook["Risk Cards"]
    headers = [str(cell.value).lstrip("\ufeff") for cell in sheet[1]]
    assert headers == list(cards[0])
    assert sheet.max_row == len(cards) + 1

    for row_number, card in enumerate(cards, 2):
        for column_number, key in enumerate(headers, 1):
            sheet.cell(row_number, column_number).value = excel_value(key, card[key])

    validation = workbook["Validation"]
    for row in validation.iter_rows():
        if row[0].value == "Active L4 cards":
            row[1].value = 622
        elif row[0].value == "L3 categories":
            row[1].value = 55
        elif row[0].value == "Unique L4 IDs":
            row[1].value = "622 of 622"

    workbook.save(TEMP)
    TEMP.replace(WORKBOOK)
    print("Synchronized HEART_L4_Risk_Cards.xlsx: 622 L4 cards and 55 L3 categories.")


if __name__ == "__main__":
    main()
